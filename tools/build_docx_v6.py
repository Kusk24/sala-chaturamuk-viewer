# Paper Version 6, the final version. Same pattern as every earlier version: the full
# paper of the previous version, with this round's progress added IN PLACE - a sentence
# block on the abstract, an update paragraph in the introduction, capture and method
# paragraphs in Section III, results subsections in Section IV-D, and a closing
# paragraph in the conclusion. New figures and tables are inserted after the last
# existing ones so the auto-numbered Fig. 1-19 and Tables I-IV keep their numbers.
# Formatting is cloned from Version 5's own paragraphs and Table IV.
import copy, os, sys
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

PAPER = '/Users/kusk/Desktop/Computer Vision/Term-Project/sala-chaturamuk-viewer/paper'
FIGS  = f'{PAPER}/figs_v6'
doc = Document(f'{PAPER}/CSX4213_Sala_Thai_Version5.docx')
P = list(doc.paragraphs)
TABLE_TEMPLATE = doc.tables[3]._tbl
USE_RUNS_FIG = False   # the Sep 10 comparison renders are illegible at column width; Table V carries it

# Version 5 paragraphs used as anchors / formatting templates (indices checked by text below)
ANCHORS = {7: 'Abstract', 12: 'This version is the third update', 25: 'A third capture',
           61: 'Gaussian splatting from the September', 62: 'Camera poses were recovered',
           63: 'One defect stopped', 98: '', 99: 'Reconstruction density', 112: 'Gaussian splatting, August',
           113: '*15,000', 117: 'Where the score', 118: 'Recomputing PSNR', 135: 'The September model is published',
           136: 'Conclusion', 138: 'Since Version 4', 152: 'R. Zhang'}
for i, start in ANCHORS.items():
    assert P[i].text.strip().startswith(start), (i, P[i].text[:60])

def set_text(p_el, text):
    runs = p_el.findall(qn('w:r'))
    best = max(runs, key=lambda r: len(''.join(t.text or '' for t in r.iter(qn('w:t')))), default=None)
    rpr = best.find(qn('w:rPr')) if best is not None else None
    rpr = copy.deepcopy(rpr) if rpr is not None else None
    for ch in list(p_el):
        if ch.tag != qn('w:pPr'):
            p_el.remove(ch)
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); p_el.append(r)
    return p_el

def clone(idx, text):
    return set_text(copy.deepcopy(P[idx]._p), text)

class Cursor:
    def __init__(self, idx): self.el = P[idx]._p
    def put(self, el): self.el.addnext(el); self.el = el; return el
    def para(self, tpl, text): return self.put(clone(tpl, text))
    def figure(self, fname, caption, width):
        img = copy.deepcopy(P[98]._p)
        for ch in list(img):
            if ch.tag != qn('w:pPr'):
                img.remove(ch)
        self.put(img)
        Paragraph(img, doc._body).add_run().add_picture(f'{FIGS}/{fname}', width=Inches(width))
        self.put(clone(99, caption))
    def table(self, title, subheads, rows, widths):
        t = copy.deepcopy(TABLE_TEMPLATE)
        trs = t.findall(qn('w:tr')); row0, row1, tpl = trs[0], trs[1], trs[2]
        for tr in trs[2:]:
            t.remove(tr)
        grid = t.find(qn('w:tblGrid'))
        if len(widths) == 4:                                  # drop one data column
            grid.remove(grid.findall(qn('w:gridCol'))[-1])
            for tr in (row1, tpl):
                tr.remove(tr.findall(qn('w:tc'))[-1])
        for g, w in zip(grid.findall(qn('w:gridCol')), widths):
            g.set(qn('w:w'), str(w))
        title_tc = row0.findall(qn('w:tc'))[1]
        title_tc.find(qn('w:tcPr')).find(qn('w:gridSpan')).set(qn('w:val'), str(len(widths) - 1))
        title_tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(sum(widths[1:])))
        set_text(title_tc.find(qn('w:p')), title)
        for tr in (row1, tpl):
            for tc, w in zip(tr.findall(qn('w:tc')), widths):
                tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
        for tc, s in zip(row1.findall(qn('w:tc'))[1:], subheads):
            set_text(tc.find(qn('w:p')), s)
        for n, values in enumerate(rows, 1):
            tr = copy.deepcopy(tpl)
            for tc, v in zip(tr.findall(qn('w:tc')), [str(n)] + list(values)):
                set_text(tc.find(qn('w:p')), v)
            t.append(tr)
        return self.put(t)

# ---------------------------------------------------------------- figure numbering
order = (['runs'] if USE_RUNS_FIG else []) + ['photos', 'heldout', 'clean', 'coverage', 'above', 'turn']
FIG = {k: 20 + i for i, k in enumerate(order)}

# ---------------------------------------------------------------- abstract
abs_p = P[7]._p
last = abs_p.findall(qn('w:r'))[-1]
r = copy.deepcopy(last)
for tnode in r.findall(qn('w:t')):
    r.remove(tnode)
tnode = OxmlElement('w:t'); tnode.set(qn('xml:space'), 'preserve')
tnode.text = (' The fourth and final update covers the last week. Two further runs on the September '
    'photographs first removed the colour sky mask, which raised the held-out score from 20.10 dB to '
    '23.45 dB, and then replaced it with a learned sky segmentation excluded from the loss, which scored '
    '22.61 dB on the pavilion pixels against 22.34 dB without it and removed almost every Gaussian from '
    'the sky. Neither run could recover the upper surface of the roof, which no photograph had seen. We therefore '
    'photographed a smaller model of the sala indoors, in six rings that include views looking down on '
    'the roof. All 301 photographs registered into a single model, which scores 25.08 dB PSNR and 0.851 '
    'SSIM on 38 held-out views and, unlike the full-size pavilion, has a roof that was photographed from '
    'above. Comparing the two shows that the gaps in the full-size reconstruction came from where the '
    'camera could stand, not from the reconstruction method.')
r.append(tnode); abs_p.append(r)

# ---------------------------------------------------------------- introduction
c = Cursor(12)
c.para(12, 'This version is the fourth and final update and covers the week since Version 5. We made '
    'two more runs on the September pavilion photographs to deal with the sky mask, which Version 5 had '
    'measured as the largest avoidable loss. We also accepted that the roof of the full-size pavilion '
    'was not going to be photographed with the access and equipment we had, since three captures had now '
    'missed it. Instead we photographed a smaller model of the same sala indoors, where a camera can be '
    'held above the roof as easily as beside it, and reconstructed it with the same pipeline. The model '
    'is reported next to the full-size pavilion (Section IV-D) so that the effect of viewpoint coverage '
    'can be seen on the same building form. As in earlier versions, the image-based renderer is '
    'unchanged and remains the contribution of the work.')

# ---------------------------------------------------------------- methodology: capture
c = Cursor(25)
c.para(25, 'A fourth capture was made on 12 September 2026 of a smaller model of the sala: a carved, '
    'gilded model standing on a wooden board on a white cloth, photographed indoors. We used one phone '
    '(iPhone 17 Pro) at one focal length, 24 mm equivalent on all 301 frames, in a single 28-minute '
    'session, and kept the EXIF. The model was photographed in six rings of 41 to 65 frames at different '
    'heights, and the later rings look down onto the roof (Fig. %d). The model was not moved during the '
    'capture. Over the first twenty frames, the image corners, which show only background, change by a '
    'mean of 13.88 grey levels between consecutive photographs, so the camera moved around a still object '
    'rather than the object turning in front of a fixed camera, which structure from motion could not '
    'have used. The HEIC files were converted to 1600-pixel JPEG as before. The phone had stored '
    'landscape pixels together with an EXIF rotation flag; COLMAP reads the pixels and ignores the flag, '
    'while most image software applies it, so the rotation was written into the pixels before any '
    'processing and all 301 frames are 1200 x 1600.' % FIG['photos'])

# ---------------------------------------------------------------- methodology: reconstruction
c = Cursor(63)
c.para(61, 'Two further runs on the September pavilion')
c.para(62, 'Both runs reused the COLMAP result of Version 5 unchanged, with 395 registered photographs '
    'and the same 50 held out, and trained for 30,000 iterations, so they differ only in how the sky is '
    'treated. The first run dropped the colour-based sky mask. The photographs were used as taken, the '
    'model was allowed to place Gaussians in the sky, and the background was removed afterwards by '
    'cropping about the true vertical of the scene and discarding oversized Gaussians. The crop radius '
    'was also widened, because the one used in Version 5 had cut away part of the terrace. The second run '
    'replaced the colour rule with a semantic segmentation network, SegFormer-B2 [12] trained on ADE20K '
    '[13], applied to every frame at 1024 pixels on the long edge. The sky it labels is passed to '
    'training as an alpha channel and the rendered image is multiplied by it before the loss, so masked '
    'pixels give no gradient while the photographs themselves stay unaltered. The mask was eroded 6 '
    'pixels into the sky, and sparse points seen mostly as sky were removed before training, 1,409 of '
    '240,357.')
c.para(61, 'The sala model')
c.para(62, 'The model went through the same notebook with three changes. The first is matching. Every '
    'pavilion run used sequential matching with loop detection off because of the four-fold symmetry '
    '(Section III-B), but that setting does not suit six separate rings, so we checked it on the prepared '
    'photographs before running COLMAP. Neighbouring frames in a ring share 241 RANSAC inliers. The last '
    'frame of one ring and the first of the next share only 12 to 25, and the strongest links between '
    'rings, 78 to 91 inliers, are between photographs about 60 frames apart in capture order, which no '
    'sequential window reaches. A pair of views 90 degrees apart in the same ring shares only 9 inliers, '
    'so the symmetry problem did not appear: the wood grain and the folds of the cloth look different from '
    'every direction. We therefore used exhaustive matching over all 45,150 image pairs. The second '
    'change is that no sky mask was used, since there is no sky indoors. The third concerns the published '
    'model. The notebook’s floater filter had been tuned on the outdoor scenes and removed very '
    'little here (Section IV-D), so the trained Gaussians were cleaned afterwards with three geometric '
    'cuts, in the reconstruction’s own units where the camera ring has a radius of 3.67: everything '
    'more than 0.90 from the vertical axis of the model, everything below the top of the board or above '
    'the top of the model, and every Gaussian with a largest axis over 0.05 or an opacity under 0.15. A '
    'gamma of 1.9 was applied to the colours because the object is dark against the viewer’s black '
    'background. None of this retrains the model, and the held-out scores in Section IV-D belong to the '
    'model before cleaning.')

# ---------------------------------------------------------------- results
c = Cursor(135)
c.para(117, 'The pavilion after Version 5')
c.para(112, 'Pavilion runs on the September photographs')
c.table('Sky Handling, Same 395 Photographs and 50 Held-Out Views', ['Measure', 'Version 5', 'Sky trained', 'Sky masked'], [
    ['Sky treatment', 'colour mask', 'none, cropped', 'learned mask'],
    ['Trained', '7 Sep.', '9 Sep.', '10 Sep.'],
    ['Whole-frame PSNR', '20.10 dB', '23.45 dB', '14.45 dB'],
    ['Pavilion-only PSNR', '–', '22.34 dB', '22.61 dB'],
    ['SSIM', '0.739', '0.769', '0.683'],
    ['LPIPS', '0.342', '0.333', '0.403'],
    ['Worst view', '10.45 dB', '14.78 dB', '9.58 dB'],
    ['Gaussians published', '291,764', '593,708', '597,067'],
    ['Gaussians in sky*', '–', '632', '34']], [430, 1400, 950, 1000, 950])
c.para(113, '*Gaussians that lie in sky in at least half of the held-out views that see them. The '
    'whole-frame score of the sky-masked run is low by design: that model renders no sky, while the '
    'held-out photographs are 28.3% sky on average.')
runs_ref = (' (Fig. %d)' % FIG['runs']) if USE_RUNS_FIG else ''
c.para(118, 'Table V sets the two runs beside the Version 5 model. Removing the colour mask helped straight '
    'away. On the same photographs and held-out views, the whole-frame score rose from 20.10 dB to 23.45 '
    'dB and the worst view from 10.45 dB to 14.78 dB, more than the 2.2–2.4 dB that Version 5 had '
    'attributed to the mask; the wider crop contributes to this as well. The run with the learned mask has '
    'a lower whole-frame score, 14.45 dB, but that number measures the empty sky the model was trained to '
    'leave out, in photographs that are between 1.2% and 43.5% sky. Scored only on the pixels of the '
    'pavilion, with the same masks and views for both runs, it reaches 22.61 dB against 22.34 dB and is '
    'better in 42 of the 50 views. The bigger change is outside the building: Gaussians sitting in the '
    'sky in at least half of the views that see them fell from 632 to 34. Some haze along the roofline '
    'remains in both runs, because those Gaussians lie on the silhouette of the building and are seen '
    'against sky from only a few directions' + runs_ref + '. Neither run changes the measurement of '
    'Version 5: no camera was above the roof, so its upper surface was never observed in either run.')
if USE_RUNS_FIG:
    c.figure('fig_v6_pavilion_runs.png', 'RUNS_CAPTION', 3.3)

c.para(117, 'The sala model')
c.figure('fig_v6_photos.jpg', 'The two subjects as photographed: (a) the full-size pavilion from its '
    'terrace loop, (b) the smaller model of the sala from the front at eye level and from above.', 3.0)
c.para(118, 'All 301 photographs registered into a single model with no second sub-model, which never '
    'happened with a capture of the full-size pavilion. The run took 7 h 27 min on Kaggle, including '
    'exhaustive matching and 30,000 iterations of training. On 38 held-out photographs the model scores '
    '25.08 dB PSNR, 0.851 SSIM and 0.371 LPIPS, with individual views between 16.58 dB and 30.58 dB and a '
    'median of 25.93 dB. The best views reproduce the carving, the coloured inlay and the red base closely '
    '(Fig. %d). These are the highest PSNR and SSIM in the project, but they need care. The model is a '
    'small object under even indoor light in front of a plain backdrop, which is an easier target than a '
    'building outdoors, and its LPIPS is worse than the pavilion’s (0.371 against 0.333). The scores '
    'show that the capture was good; they do not show that the method improved.' % FIG['heldout'])
c.figure('fig_v6_model_heldout.jpg', 'Three of the 38 held-out photographs of the model (top) beside '
    'renderings from the same viewpoints (bottom). The lowest-scoring view, right, is covered by the haze '
    'described in the text.', 3.3)
c.para(118, 'The trained model was surrounded by haze. The white cloth has almost no texture, so nothing in '
    'training pins down the empty space between the camera and the object, and large transparent Gaussians '
    'there cost the loss very little. It is the same problem as the sky outdoors. The notebook’s '
    'floater filter did not help: its threshold, 0.04 of the ring radius or 0.147 units, is larger than '
    '99.9% of the Gaussians in this scene, whose 99.9th-percentile largest axis is 0.143, so it removed '
    'only 3,557 of 225,756. After the cuts described in Section III-C, 77,626 of the 222,199 Gaussians '
    'remain. Rendered offline from one fixed camera, the notebook’s own opening view, 93.1% of the '
    'frame is not background in the trained model and 56.0% of it is low in saturation, which is the '
    'veil in front of the object with the board beneath it. After cleaning, 8.8% of the frame is not '
    'background and 1.3% is grey, so what is left is the object (Fig. ' + str(FIG['clean']) + '). The '
    'viewer page uses the cleaned model and states that the scores above belong to the model before '
    'cleaning.')
c.figure('fig_v6_model_clean.jpg', 'The model rendered from one fixed camera (a) as trained and (b) after '
    'the geometric cuts and colour lift. Non-background pixels fall from 93.1% of the frame to 8.8%.', 3.3)

c.para(117, 'The full-size pavilion and the model compared')
c.para(112, 'The full-size pavilion and the sala model')
c.table('Coverage and Results of the Two Captures', ['Measure', 'Full-size', 'Model'], [
    ['Setting', 'outdoors, by the lake', 'indoors, on a table'],
    ['Photos / registered', '446 / 395', '301 / 301'],
    ['Matching', 'sequential', 'exhaustive'],
    ['Highest camera', '−9.7°', '+67.7°'],
    ['Cameras above horizontal', '0', '237'],
    ['Cameras above subject top', '0', '151'],
    ['Highest camera vs. top', '−0.46 ring', '+0.70 ring'],
    ['Surface radius / ring', '0.92', '0.20'],
    ['30° sectors covered', '12 of 12', '12 of 12'],
    ['PSNR / SSIM', '23.45 / 0.769', '25.08 / 0.851'],
    ['LPIPS', '0.333', '0.371'],
    ['Gaussians in viewer', '593,708', '77,626']], [430, 1700, 1300, 1300])
c.para(113, 'Pavilion scores are from the sky-trained run. Elevation is measured about the centre of each '
    'model with the world up-vector taken from the cameras’ own up axes, as in Fig. 17. Heights are '
    'measured from the top of the reconstructed subject and, like the surface radius (90th percentile), '
    'in units of the camera-ring radius, the median horizontal camera distance.')
c.para(118, 'Table VI and Fig. %d show why one capture gives a model that can be turned all the way round '
    'and seen from above while the other does not. Both captures went all the way round, with at least 20 '
    'photographs in every 30-degree sector of azimuth. The difference is height and distance. Every '
    'pavilion photograph was taken from below the building, the highest at −9.7 degrees about its '
    'centre, and every camera looks up at the top of the roof, by at least 25.8 degrees. For the model, 237 of the 301 cameras are above the '
    'horizontal and 151 are higher than the top of the model, looking down onto it at up to 60.6 degrees. '
    'A roof seen only from below, at a grazing angle, leaves its upper surface to guesswork. From above, '
    'the pavilion’s roof keeps its outline but is streaked and partly filled with sky colour, while '
    'the model shows both roof slopes, the ridge and its corner finials (Fig. %d).' % (FIG['coverage'], FIG['above']))
c.figure('fig_v6_coverage.png', 'Camera coverage of the two captures. (a) Elevation of every registered '
    'camera about the centre of its model. (b) Camera positions in a vertical slice: height above the top '
    'of the subject against horizontal distance, both in camera-ring radii, with each reconstructed '
    'surface shaded. No pavilion camera is above its roof.', 3.3)
c.para(118, 'Distance matters for a similar reason. The pavilion’s reconstructed surface, building and '
    'terrace together, reaches 0.92 of the camera-ring radius, so the cameras were walking along its edge, '
    'because the loops used for the model were taken close in; the 171 photographs taken from further '
    'away were never trained. The model reaches only 0.20 of its ring, so every photograph held the whole '
    'object and a short movement of the phone covered a wide range of directions. The background also '
    'differs. Behind the pavilion there is sky, 28.3% of each held-out photograph on average, and trees '
    'and buildings far enough away that walking round the pavilion gives them almost no parallax. Behind '
    'the model there is a board and a cloth close enough to change with every step, and this is what broke '
    'the four-fold symmetry that had forced sequential matching on the pavilion.')
c.figure('fig_v6_from_above.jpg', 'Both reconstructions rendered from directly above with the same '
    'renderer: (a) the full-size pavilion, whose roof keeps its outline but is streaked and tinted with sky '
    'colour because it was only seen from below, and (b) the model, whose roof was photographed from '
    'above.', 3.3)
c.para(118, 'The model does not replace the building. It has no underside, because it stood on a board, '
    'and it is a smaller copy rather than the pavilion itself. What it shows is that the gaps in the '
    'full-size reconstruction come from where the camera could be placed and not from Gaussian splatting '
    'or from the amount of training: with views from above and from far enough away, the same pipeline '
    'produces a model that holds together from every side (Fig. %d).' % FIG['turn'])
c.figure('fig_v6_turntable_2x2.jpg', 'The cleaned model rendered offline at 0, 90, 180 and 270 degrees of '
    'azimuth, 20 degrees above the horizontal.', 3.0)

# ---------------------------------------------------------------- conclusion
c = Cursor(138)
c.para(138, 'This is the final version of the project. In the last week we ran the pavilion twice more and '
    'photographed a smaller model of the sala. Removing the colour sky mask raised the pavilion’s '
    'held-out score from 20.10 dB to 23.45 dB, and a learned sky mask then improved the pavilion pixels '
    'slightly, to 22.61 dB against 22.34 dB, while removing nearly all Gaussians from the sky. Neither run '
    'could fill in the upper surface of the roof, because no photograph had seen it. The model, captured with cameras around '
    'and above it, registered all 301 photographs and is the first reconstruction of the sala whose roof '
    'was actually photographed, at 25.08 dB on held-out views. Side by side, the two show that what '
    'limited the full-size result was the capture, meaning where a camera could and could not be placed, '
    'rather than Gaussian splatting or the amount of training. Anyone continuing the work should start '
    'with elevated photographs of the real pavilion and apply the same geometric cleaning, with a floater '
    'threshold set from each scene rather than carried over from another. The image-based viewer remains '
    'the contribution of this work, and the reconstructions are reported beside it as a comparison.')

# ---------------------------------------------------------------- references
c = Cursor(152)
c.para(152, 'E. Xie, W. Wang, Z. Yu, A. Anandkumar, J. M. Alvarez, and P. Luo, “SegFormer: Simple and '
    'efficient design for semantic segmentation with transformers,” in Adv. Neural Inf. Process. '
    'Syst. (NeurIPS), vol. 34, 2021, pp. 12077–12090.')
c.para(152, 'B. Zhou, H. Zhao, X. Puig, S. Fidler, A. Barriuso, and A. Torralba, “Scene parsing through '
    'ADE20K dataset,” in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), Honolulu, HI, USA, '
    '2017, pp. 633–641.')

RUNS_CAPTION = os.environ.get('RUNS_CAPTION', '')
if USE_RUNS_FIG:
    assert RUNS_CAPTION, 'set RUNS_CAPTION'
    for p in doc.paragraphs:
        if p.text == 'RUNS_CAPTION':
            set_text(p._p, RUNS_CAPTION)

DST = f'{PAPER}/CSX4213_Sala_Thai_Version6.docx'
doc.save(DST)
print('saved', DST, round(os.path.getsize(DST) / 1e6, 1), 'MB; figures', FIG)
